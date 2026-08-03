import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "output" / "docx"
OUT.mkdir(parents=True, exist_ok=True)

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
HEADER_FILL = "E8EEF5"
LIGHT_FILL = "F4F6F9"
RISK_FILL = "FFF4E5"
OK_FILL = "EAF7EA"
TEXT = RGBColor(31, 41, 55)
MUTED = RGBColor(71, 85, 105)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(width)


def set_run(run, size=None, bold=None, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_para(doc, text="", style=None, bold=False, size=None, color=None, after=6):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if text:
        r = p.add_run(text)
        set_run(r, size=size, bold=bold, color=color or TEXT)
    return p


def add_title(doc, title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run(r, size=22, bold=True, color=RGBColor(15, 23, 42))
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(subtitle)
    set_run(r, size=11, color=MUTED)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("DRAFT - NOT EFFECTIVE UNTIL REVIEWED, APPROVED, DATED, AND SIGNED")
    set_run(r, size=9, bold=True, color=RGBColor(155, 28, 28))


def add_h1(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    set_run(r, size=16, bold=True, color=RGBColor(46, 116, 181))
    return p


def add_h2(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    set_run(r, size=13, bold=True, color=RGBColor(46, 116, 181))
    return p


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(item)
        set_run(r, size=10.5, color=TEXT)


def add_table(doc, rows, widths, header=True, fills=None):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    set_table_width(table, widths)
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if header and r_idx == 0:
                set_cell_shading(cell, HEADER_FILL)
            elif fills and r_idx in fills:
                set_cell_shading(cell, fills[r_idx])
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(str(value))
            set_run(run, size=9.3, bold=(header and r_idx == 0), color=TEXT)
    return table


def setup_doc(title_left, title_right):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    header = section.header.paragraphs[0]
    header.text = ""
    r = header.add_run(f"{title_left} | {title_right}")
    set_run(r, size=8, bold=True, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Draft packet - page ")
    set_run(r, size=8, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    return doc


def build_board_packet():
    doc = setup_doc("VCC Board Action Packet", "DOCS 8404 Launch Readiness")
    add_title(
        doc,
        "VCC Board Action Packet",
        "Governance, ownership/control, compliance, and transportation launch approval framework",
    )
    add_para(
        doc,
        "Prepared for VILIGANS COMMAND CORPORATION as an internal draft packet. This packet prepares decisions and evidence; it does not execute filings, issue stock, bind insurance, submit grants, sign contracts, or authorize paid passenger service.",
        after=8,
    )
    add_h1(doc, "Executive Decision")
    add_table(
        doc,
        [
            ["Decision Area", "Recommended Board Position", "Reason"],
            ["External control lane", "Use Alawndus Davis -> VCC -> DOCS 8404 for active launch work.", "This is the cleanest lane for banking, SAM, VetCert readiness, fleet, funding, and partner-facing materials."],
            ["Unverified ownership language", "Do not use externally unless counsel/CPA approves a limited internal purpose.", "Prior binder materials contained inconsistent ownership narratives that can slow banks, grant reviewers, SBA, Medicaid, and insurers."],
            ["Outreach", "Prepare only until evidence, authority, insurance, and launch gates are complete or waived.", "Prevents unsupported claims and premature commitments."],
        ],
        [1.65, 2.4, 2.45],
    )

    doc.add_page_break()
    add_h1(doc, "Source-Proof Status")
    add_table(
        doc,
        [
            ["Requirement", "Current Evidence", "Status", "Next Action"],
            ["VCC Wyoming corporation", "Filed Articles, Filing ID 2025-001851867, and July 28, 2026 certificate of good standing found in Google Drive.", "Verified", "Insert into final corporate binder tab."],
            ["Authorized shares", "Filing info shows 1,000,000 common shares, $0.001 par value, 0 preferred shares.", "Verified", "Use these values in stock issuance records."],
            ["Bylaws", "Generic/blank bylaws PDF found.", "Template only", "Complete VCC-specific bylaws and approve."],
            ["Initial resolutions", "Generic/blank initial resolutions PDF found.", "Template only", "Complete VCC-specific resolutions and approve."],
            ["DOCS 8404 state filing", "Operating agreement found; official Wyoming LLC filing proof not found in visible files.", "Open", "Pull official SOS record or file/verify LLC."],
            ["SAM/UEI and VetCert", "Readiness drafts exist; active registration and certification approval not verified.", "Open", "Complete SAM, Grants.gov, and VetCert readiness evidence."],
        ],
        [1.5, 2.45, 0.95, 1.6],
        fills={1: OK_FILL, 2: OK_FILL, 3: RISK_FILL, 4: RISK_FILL, 5: RISK_FILL, 6: RISK_FILL},
    )

    doc.add_page_break()
    add_h1(doc, "Draft Resolutions")
    resolutions = [
        ("Resolution 001 - Active Control Structure", "The corporation confirms the active external-facing launch structure as Alawndus Davis -> VILIGANS COMMAND CORPORATION -> DOCS 8404 Logistics & NEMT LLC, subject to verification of DOCS 8404 state records and counsel/CPA review where required."),
        ("Resolution 002 - Governance Completion", "The corporation directs completion of VCC-specific bylaws, initial board consent, officer appointments, signing authority, and conflict/related-party controls before external submissions."),
        ("Resolution 003 - Stock Issuance Readiness", "The corporation authorizes preparation of a founder stock issuance package using the filed share authorization. No third-party securities offering, investor solicitation, pledge, or transfer is authorized by this draft."),
        ("Resolution 004 - Federal Readiness", "The corporation authorizes preparation of SAM.gov, Grants.gov, and SBA VetCert readiness records. No federal representation or application is authorized until reviewed and approved."),
        ("Resolution 005 - Transportation Launch Gate", "The corporation prohibits paid passenger service until WYDOT/FMCSA path, insurance, vehicle readiness, driver files, operating policies, and incident/privacy controls are complete and approved."),
        ("Resolution 006 - Partner And Funding Outreach", "The corporation authorizes drafting of outreach, MOU, sponsorship, and grant materials. No contract, MOU, sponsorship, grant submission, or financial obligation is effective without approval and signature."),
    ]
    for title, body in resolutions:
        add_h2(doc, title)
        add_para(doc, body)

    doc.add_page_break()
    add_h1(doc, "Launch Gate Checklist")
    add_table(
        doc,
        [
            ["Gate", "Go Standard", "Current Status"],
            ["Entity records", "VCC and DOCS 8404 source-proof records complete.", "VCC verified; DOCS 8404 open."],
            ["Governance", "Bylaws, officer appointments, stock ledger, cap table, and signing authority approved.", "Open."],
            ["Authority", "WYDOT/FMCSA path documented; application filed if required.", "Open."],
            ["Insurance", "Commercial passenger/NEMT coverage bound.", "Open."],
            ["Vehicle", "Accessible passenger vehicle titled/leased, inspected, and insured.", "Open."],
            ["Staffing", "Driver and backup driver files complete.", "Open."],
            ["Partner", "At least one LOI, MOU, paid pilot agreement, or sponsor commitment secured.", "Open."],
        ],
        [1.55, 3.35, 1.6],
    )

    add_h1(doc, "Signature Blocks")
    add_para(doc, "These signature lines are placeholders. Do not sign until the final packet is reviewed and approved.")
    add_table(
        doc,
        [
            ["Role", "Name", "Signature", "Date"],
            ["President / CEO", "Alawndus Davis", "", ""],
            ["Secretary", "", "", ""],
            ["Director", "", "", ""],
        ],
        [1.45, 1.7, 2.1, 1.25],
    )
    path = OUT / "VCC_Board_Action_Packet_DRAFT.docx"
    doc.save(path)
    return path


def build_partner_packet():
    doc = setup_doc("VCC Partner Outreach Packet", "Draft MOU And Scripts")
    add_title(
        doc,
        "VCC Partner Outreach Packet",
        "Anchor partner, sponsor, VSO, healthcare, county, and disability-provider outreach drafts",
    )
    add_para(
        doc,
        "This packet contains external-facing draft language for review. It should not be sent until VCC approves the service area, claims, authority status, insurance status, vehicle readiness, and escalation rules.",
    )
    add_h1(doc, "Controlled Outreach Rule")
    add_table(
        doc,
        [
            ["Rule", "Approved Draft Position"],
            ["Certification", "Do not claim SBA SDVOSB, Medicaid provider, WYDOT-authorized carrier, or fully insured carrier status until verified."],
            ["Service", "Describe the pilot as scheduled, non-emergency, non-ambulance transportation only."],
            ["Commitments", "Do not promise trips, dates, prices, Medicaid billing, wheelchair transport, or grant awards until approved."],
            ["Data", "Collect only minimum necessary referral and trip information."],
        ],
        [1.45, 5.05],
    )
    add_h1(doc, "Anchor Partner Offer")
    add_para(doc, "The recommended first revenue offer is a 90-day pilot with a fixed monthly readiness fee, per-trip charge, limited service area, weekly reporting, and a documented launch gate.")
    add_table(
        doc,
        [
            ["Offer Element", "Draft Position"],
            ["Pilot period", "90 days after authority, insurance, vehicle, driver, and policy gates are complete."],
            ["Payment model", "Monthly readiness fee plus per-trip charge, or sponsor-funded ride block."],
            ["Reporting", "Weekly trips, miles, denials, no-shows, on-time rate, rider need, and partner source."],
            ["Limitations", "No emergency transport, no medical treatment, no Medicaid billing until enrollment/billing path is verified."],
        ],
        [1.6, 4.9],
    )
    add_h1(doc, "Email Script - Healthcare Partner")
    add_para(doc, "Subject: Wyoming rural transportation pilot for patient access", bold=True)
    add_para(doc, "Hello [Name],")
    add_para(doc, "VILIGANS COMMAND CORPORATION is preparing a controlled rural transportation pilot through DOCS 8404 Logistics & NEMT LLC to support scheduled, non-emergency access for Wyoming residents who struggle to reach appointments and essential services.")
    add_para(doc, "We are looking for one anchor healthcare partner to help validate the pilot around recurring access needs such as primary care, dialysis, rehab, behavioral health, pharmacy access, and follow-up appointments.")
    add_para(doc, "The pilot would be limited, scheduled, insured, documented, and launched only after operating authority, vehicle, driver, and safety gates are complete. We are not proposing ambulance or emergency medical transport.")
    add_para(doc, "Could we schedule a 20-minute call to discuss whether your team sees unmet transportation needs and whether a letter of support, referral workflow, sponsored ride block, or small paid pilot would be useful?")
    add_h1(doc, "Email Script - VSO / Veteran Partner")
    add_para(doc, "Subject: Veteran transportation operator partnership discussion", bold=True)
    add_para(doc, "Hello [Name],")
    add_para(doc, "VILIGANS COMMAND CORPORATION is preparing DOCS 8404 Logistics & NEMT LLC as a Wyoming rural transportation operator with a veteran-led mission and SBA VetCert readiness work underway.")
    add_para(doc, "We understand that some veteran transportation grants require a qualified VSO or State Veterans Service Agency lead applicant. VCC is not assuming direct eligibility where a qualified VSO/SVSA is required. Instead, we are preparing an operator/subrecipient packet for eligible partners that need rural transportation capacity.")
    add_para(doc, "Could we discuss whether your organization has veteran ride needs, grant-readiness needs, or interest in a future MOU or letter of support?")
    doc.add_page_break()
    add_h1(doc, "MOU Skeleton")
    add_table(
        doc,
        [
            ["Section", "Draft Content"],
            ["Purpose", "Explore a controlled Wyoming rural mobility pilot for scheduled, non-emergency transportation."],
            ["Pilot scope", "County/service area, pilot period, trip types, rider groups, operating days, accessibility capacity, and payment route."],
            ["VCC/DOCS duties", "Maintain required authority, insurance, driver files, vehicle records, policies, dispatch, incident reporting, and trip data."],
            ["Partner duties", "Identify referral categories, referral contacts, funding route, feedback process, and letter-of-support possibilities."],
            ["Non-binding status", "No payment, trips, grant submission, or contract obligation unless separately approved and signed."],
        ],
        [1.45, 5.05],
    )
    add_h1(doc, "Phone Script")
    add_para(doc, "Hello, this is Alawndus Davis with VILIGANS COMMAND CORPORATION. We are preparing a controlled Wyoming rural transportation pilot for scheduled, non-emergency rides serving seniors, veterans, individuals with disabilities, and residents who have trouble reaching healthcare or essential services.")
    add_numbered(
        doc,
        [
            "Are you seeing unmet transportation needs in your community or patient/client population?",
            "Would your organization consider a planning call, letter of support, referral partnership, sponsored ride block, or pilot MOU?",
            "Who is the best person to speak with about transportation access, community health, or partnership opportunities?",
        ],
    )
    path = OUT / "VCC_Partner_Outreach_Packet_DRAFT.docx"
    doc.save(path)
    return path


def read_csv_rows(name):
    path = ROOT / "foundation" / "launch_package" / name
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.reader(handle))


def build_vehicle_funding_close_packet():
    doc = setup_doc("VCC Vehicle Funding Close Packet", "Wheatland August 2026 Sprint")
    add_title(
        doc,
        "Vehicle Funding Close Packet",
        "Wheatland pilot funding, vehicle, authority, insurance, and approval controls",
    )
    add_para(
        doc,
        "Prepared for VILIGANS COMMAND CORPORATION as an internal and review-ready draft packet. This packet supports quote, fit, timing, sponsor, and readiness discussions. It does not authorize sending, signing, applications, credit pulls, deposits, insurance binding, vehicle commitments, filings, invoices, restricted funds, or paid passenger service.",
        after=8,
    )

    add_h1(doc, "Executive Target")
    add_table(
        doc,
        [
            ["Target", "Current position", "Evidence needed"],
            ["Funding by August 9, 2026", "Realistic for commitments, bridge support, sponsor/readiness pledges, or preapproval; not realistic for grant cash.", "Signed pledge, payment proof, approved bridge memo, lender preapproval, or written term sheet."],
            ["Pilot vehicle by August 15, 2026", "Possible only through an immediately available accessible van purchase/lease, short-term rental bridge, or approved partner vehicle.", "Written quote or rental confirmation, insurance bindability, WYDOT path, approval memo, and funding source."],
            ["Paid passenger launch", "Not approved by this packet.", "Authority, insurance, vehicle, driver, policy, privacy, rate-card, and partner gates complete."],
        ],
        [1.45, 2.55, 2.5],
        fills={1: RISK_FILL, 2: RISK_FILL, 3: RISK_FILL},
    )

    add_h1(doc, "Approval Control")
    add_table(
        doc,
        [
            ["Allowed after CEO approval", "Not allowed without separate approval"],
            ["Controlled quote, fit, and timing contact with listed vendors, lenders, insurers, WYDOT, and sponsor/anchor targets.", "Applications, credit pulls, deposits, vehicle reservations, purchase/lease/rental agreements, insurance bind orders, authority filings, invoices, restricted funds, sponsor logo promises, public launch claims, or paid passenger service."],
            ["Response logging in 49_External_Response_Evidence_Log_2026-08-02.csv and 52_Vehicle_Funding_Quote_Intake_Scorecard_2026-08-02.csv.", "Treating a conversation, public listing, or draft as funding proof or vehicle readiness proof."],
        ],
        [3.15, 3.35],
        fills={1: RISK_FILL, 2: RISK_FILL},
    )

    add_h1(doc, "Minimum Funding Cash Triggers")
    add_para(
        doc,
        "Use 61_Minimum_Funding_Stack_And_Cash_Trigger_Memo_2026-08-02.md before calling the August 9 funding sprint closed. Public listings and rate pages support triage only; written proof and approval are still required.",
    )
    add_table(
        doc,
        [
            ["Vehicle path", "Count by August 9 only if", "Internal cash trigger"],
            ["Rental bridge", "Written availability, total due, business/demo-use rules, insurance requirements, mileage, pickup/delivery, cancellation terms, and approval path.", "$5,000-$8,000 available or authorized"],
            ["Lease start", "Written lease/preapproval response with vehicle class or VIN, total due at delivery, term, mileage, guaranty, fees, and delivery timing.", "$8,000-$15,000 available or authorized"],
            ["Financed purchase", "Written quote, down-payment path, title/inspection/warranty status, insurance bindability, commercial-use permission, and approval sequence.", "$10,000-$20,000 bridge"],
            ["Cash purchase", "Written quote with all closing costs, insurance, registration, delivery, safety items, and reserve confirmed.", "$35,000-$45,000 close"],
            ["Partner vehicle", "Written partner capacity, authority/insurance proof, vehicle/driver controls, payment/data/liability terms, and MOU path.", "$2,500-$7,500 setup"],
        ],
        [1.25, 3.55, 1.7],
        fills={1: RISK_FILL, 2: RISK_FILL, 3: RISK_FILL, 4: RISK_FILL, 5: RISK_FILL},
    )

    add_h1(doc, "Controlled Contact Authorization")
    add_para(
        doc,
        "Use 51_Controlled_Contact_Authorization_Record_2026-08-02.md as the source-controlled approval record before any external call, email, form submission, quote request, insurance discussion, financing inquiry, or sponsor approach.",
    )
    add_table(
        doc,
        [
            ["Authorization lane", "Permitted only if approved", "Still excluded"],
            ["Authority", "Non-binding WYDOT questions on MC-100, Form E, passenger capacity, USDOT, markings, registration, and pre-authority limits.", "Filing, fee payment, Form E filing, or authority representation."],
            ["Insurance", "Eligibility, underwriting inputs, Form E capability, limits, driver/vehicle rules, and quote requirements.", "Paid application, coverage bind, or policy purchase."],
            ["Vehicle/rental", "Availability, total due, business-use rules, pickup/delivery, insurance requirements, and written quote inputs.", "Reservation, deposit, payment, lease, rental, purchase, or passenger use."],
            ["Funding", "Planning call, sponsor/readiness interest, pledge path, preliminary financing fit, and document checklist.", "Invoice, restricted funds, logo promise, credit pull, debt, guaranty, lien, or collateral pledge."],
        ],
        [1.15, 3.05, 2.3],
        fills={1: RISK_FILL, 2: RISK_FILL, 3: RISK_FILL, 4: RISK_FILL},
    )

    add_h1(doc, "Authority And Insurance Evidence")
    add_table(
        doc,
        [
            ["Lane", "Current evidence", "Sprint implication"],
            ["WYDOT authority", "WYDOT defines contract motor carrier as intrastate transportation of people or property for compensation; WYDOT lists 307-777-4850 for questions.", "Call or send the unsent authority question draft before relying on any paid passenger-service plan."],
            ["MC-100 / Form E", "WYDOT MC-100 asks passenger carriers to list passenger count; Form E is required before contract carrier authority approval if applicable.", "Insurance conversations must ask whether the carrier can issue Wyoming Form E."],
            ["Coverage threshold", "WYDOT MC-100 lists $750,000 combined single limit for contract carrier liability filings.", "Quote conversations should test at least this threshold if the authority path applies."],
            ["USDOT trigger", "WYDOT MC-100 states more than 9 passengers, including the driver, must have a USDOT number.", "Keep the first vehicle below that threshold unless VCC intentionally opens the USDOT path."],
            ["Insurance quote paths", "Insureon lists NEMT/paratransit quote support at 800-688-1984; Progressive lists NEMT quote support at 1-888-806-9598 in select states.", "Use both as eligibility and bindability checks, not as binding applications."],
        ],
        [1.2, 3.25, 2.05],
        fills={1: RISK_FILL, 2: RISK_FILL, 3: RISK_FILL, 4: RISK_FILL, 5: RISK_FILL},
    )

    add_h1(doc, "Funding Close Register")
    funding_rows = read_csv_rows("40_August_9_Funding_Close_Register_2026-08-02.csv")
    condensed_funding = [["Lane", "Target amount", "Proof required", "Status", "Next action"]]
    for row in funding_rows[1:]:
        condensed_funding.append([row[1], row[3], row[4], row[6], row[9]])
    add_table(doc, condensed_funding, [1.25, 1.0, 1.95, 1.0, 1.3])

    doc.add_page_break()
    add_h1(doc, "Quote Intake Scorecard")
    score_rows = read_csv_rows("52_Vehicle_Funding_Quote_Intake_Scorecard_2026-08-02.csv")
    condensed_score = [["Priority", "Target", "Status", "Proof needed"]]
    for row in score_rows[1:]:
        condensed_score.append([row[0], row[2], row[6], row[4]])
    add_table(doc, condensed_score, [0.55, 1.55, 1.1, 3.3])

    doc.add_page_break()
    add_h1(doc, "Public Vehicle Shortlist")
    vehicle_rows = read_csv_rows("37_Public_Vehicle_Inventory_Shortlist_2026-08-02.csv")
    condensed_vehicle = [["Priority", "Vendor", "Vehicle", "Price / terms", "Fastest use", "Next action"]]
    for row in vehicle_rows[1:]:
        condensed_vehicle.append([row[0], row[2], row[3], row[4], row[11], row[14]])
    add_table(doc, condensed_vehicle, [0.55, 1.15, 1.45, 1.15, 1.05, 1.15])

    add_h1(doc, "Sponsor / Readiness Commitment Terms")
    add_table(
        doc,
        [
            ["Term", "Draft control"],
            ["Commitment types", "Vehicle readiness sponsorship, sponsor-funded local ride block, anchor readiness fee, community access day sponsorship, in-kind support, or letter of support."],
            ["Allowed uses", "Accessible vehicle deposit/rental/lease startup, insurance deposit, inspection, securement, safety kit, registration/use documentation, dispatch/reporting setup, route planning, partner onboarding, and compliance readiness."],
            ["Recognition", "Sponsor recognition only after approval; no logo, vehicle/event placement, public statement, or service-readiness claim without review."],
            ["Reporting", "Non-PHI reporting only: ride requests, completed rides, denied/deferred rides, general service areas, cost summaries, incident-free days, and unmet-need observations."],
            ["Restrictions", "Do not treat a pledge as cash until funds are received or a binding approved written commitment exists."],
        ],
        [1.45, 5.05],
    )

    doc.add_page_break()
    add_h1(doc, "August 9 Go / No-Go Evidence")
    add_table(
        doc,
        [
            ["Evidence", "Minimum standard"],
            ["Vehicle option", "Written quote or rental confirmation with stock/VIN or rental class, conversion, entry type, securement, mileage, price, total due, timing, and commercial-use rules."],
            ["Insurance", "Broker/carrier guidance that the service model and vehicle can be insured, including required limits, exclusions, driver requirements, and Form E capability if WYDOT requires it."],
            ["Authority", "WYDOT call notes confirming MC-100, Form E, registration, marking, USDOT, and passenger-capacity implications for Wheatland paid passenger service."],
            ["Funding", "Signed pledge, readiness-fee commitment, lender preapproval, approved bridge funds, or microloan/lease path with timing and conditions."],
            ["Approval", "CEO/officer approval of any outreach, quote submission, application, credit pull, deposit, invoice, agreement, or restricted-fund acceptance."],
        ],
        [1.35, 5.15],
    )

    add_h1(doc, "CEO Approval Placeholder")
    add_para(doc, "Approval requested only for controlled external contact to gather quote, fit, and timing information.")
    add_table(
        doc,
        [
            ["Approval", "Selection"],
            ["Controlled external contact only", "Yes / No"],
            ["Excluded items or limits", ""],
            ["Authorized signer", "ALAWNDUS L. DAVIS, CEO / Founder"],
            ["Signature", ""],
            ["Date", ""],
        ],
        [2.05, 4.45],
    )

    path = OUT / "VCC_Vehicle_Funding_Close_Packet_DRAFT.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    for created in [build_board_packet(), build_partner_packet(), build_vehicle_funding_close_packet()]:
        print(created)
